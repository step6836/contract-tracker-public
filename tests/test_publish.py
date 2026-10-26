"""Exec report publishing: run with  python -m pytest tests -q"""
import json
import os
import tempfile
from datetime import date, timedelta

os.environ.setdefault("TRACKER_DATA_DIR", tempfile.mkdtemp())

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import auth, main, publish, storage

SHEETS = ["Renewals", "Transfers", "Expansions", "Leases"]  # one active-cases sheet per type, dashboard order


@pytest.fixture(autouse=True)
def config(tmp_path, monkeypatch):
    """Point config.json at a temp folder and stop debounced publishes from other tests firing mid-test."""
    monkeypatch.setattr(storage, "CONFIG_FILE", tmp_path / "config.json")
    if publish._timer:
        publish._timer.cancel()
    return tmp_path


@pytest.fixture(scope="module")
def client():
    auth.create_user("publisher", "Publisher", "correct-horse-9")
    c = TestClient(main.app)
    assert c.post("/api/login", json={"username": "publisher", "password": "correct-horse-9"}).status_code == 200
    started = (date.today() - timedelta(days=30)).isoformat()  # long enough to be stale on its first step
    r = c.post("/api/cases", json={"type": "renewal", "store_number": "77", "store_name": "Report Store",
                                   "franchisee": "Pat Example", "start_date": started,
                                   "dates": {"fa_expiration": "2099-01-01", "election_due": "2098-09-01"}})
    assert r.status_code == 200, r.text
    # waiting on the Election Form has no clock (its due date covers it), so move on to the evaluation
    assert c.post(f"/api/cases/{r.json()['id']}/steps/election_form_received/complete",
                  json={"completed_on": started}).status_code == 200
    closed = c.post("/api/cases", json={"type": "transfer", "store_number": "78", "store_name": "Closed Store",
                                        "start_date": started}).json()
    assert c.post(f"/api/cases/{closed['id']}/close", json={"outcome": "withdrawn", "reason": "Buyer walked"}).status_code == 200
    return c


def set_path(tmp_path, path):
    (tmp_path / "config.json").write_text(json.dumps({"publish_path": str(path)}))


def test_publish_creates_workbook(client, config):
    out = config / "Exec report.xlsx"
    set_path(config, out)
    assert not out.exists()
    assert publish.publish_now()["ok"]
    assert out.exists()
    assert [p.name for p in config.iterdir() if "tmp" in p.name] == []

    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", *SHEETS, "Completed & Closed"]
    assert wb.active.title == "Summary"  # execs land on the summary
    summary, active, closed = wb["Summary"], wb["Renewals"], wb["Completed & Closed"]
    for ws in (summary, *(wb[n] for n in SHEETS), closed):
        assert ws["A2"].value.startswith("Last updated: ")
        assert not ws.sheet_view.showGridLines
    assert active.freeze_panes == "A4"
    assert [c.value for c in active[3]] == publish._active_headers(signing=True)
    assert "Type" not in publish._active_headers(signing=True)              # the sheet is the type
    assert active.cell(row=3, column=1).font.bold
    assert "ActiveRenewals" in active.tables  # banded table with filter buttons
    assert active["A1"].value == "Active Renewals"
    assert "Lease Expires" in [c.value for c in wb["Leases"][3]]

    row = next(r for r in active.iter_rows(min_row=4) if r[2].value == "Report Store")
    assert "Pat Example" not in [c.value for c in row]                       # no franchisee for execs
    assert row[4].value == "Contracts"                                         # responsible team
    assert row[6].value == "Delayed" and row[6].fill.fgColor.rgb.endswith("FFEB9C")
    assert row[7].value is None or row[7].number_format == publish.DATE_FMT
    # the closed transfer isn't on the Transfers sheet, and renewals stay off it
    assert all(r[2] not in ("Report Store", "Closed Store") for r in wb["Transfers"].iter_rows(min_row=4, values_only=True))

    values = [c for r in summary.iter_rows(values_only=True) for c in r]
    for text in ("At a Glance", "Case Type", "Delayed", "On Hold", "Past Due", "Due Within 30 Days", "Closed Early",
                 "Needs Attention", "Upcoming and Past Due Signings", "Status Key", "Renewals"):
        assert text in values, text
    assert list(publish.STATUS_STYLE) == ["On Track", "Delayed", "On Hold", "Past Due"]
    assert "R-77-1" in values  # listed under Needs attention

    rows = list(closed.iter_rows(values_only=True))
    assert [list(r[:len(publish.CLOSED_HEADERS)]) for r in rows].count(publish.CLOSED_HEADERS) == 1
    done = next(r for r in rows if r[3] == "Closed Store")
    assert done[4] == "Withdrawn" and done[5] == "Buyer walked"


def test_status_precedence():
    t = date(2026, 10, 13)
    base = {"target_signing_date": "2026-11-01", "signing_date": None, "hold": None,
            "current": {"is_stale": False}}
    assert publish._status(base, t) == "On Track"
    assert publish._status({**base, "current": {"is_stale": True}}, t) == "Delayed"
    assert publish._status({**base, "current": {"is_stale": True}, "hold": {"reason": "x"}}, t) == "On Hold"
    assert publish._status({**base, "hold": {"reason": "x"}, "target_signing_date": "2026-10-01"}, t) == "Past Due"
    assert publish._status({**base, "target_signing_date": "2026-10-01", "signing_date": "2026-09-30"}, t) == "On Track"


def test_publish_replaces_existing_file(client, config):
    out = config / "report.xlsx"
    set_path(config, out)
    out.write_bytes(b"not a real workbook")  # never opened, just replaced
    assert publish.publish_now()["ok"]
    assert load_workbook(out).sheetnames == ["Summary", *SHEETS, "Completed & Closed"]


def test_missing_publish_path_does_nothing(client, config):
    (config / "config.json").write_text("{}")
    r = publish.publish_now()
    assert not r["ok"] and "publish_path" in r["detail"]
    assert [p.name for p in config.iterdir()] == ["config.json"]


def test_publish_failure_never_raises(client, config):
    set_path(config, config / "no such folder" / "report.xlsx")
    assert not publish.publish_now()["ok"]
    r = client.post("/api/publish")
    assert r.status_code == 200 and r.json()["ok"] is False


def test_publish_endpoint(client, config):
    out = config / "report.xlsx"
    set_path(config, out)
    r = client.post("/api/publish")
    assert r.status_code == 200 and r.json()["ok"], r.text
    assert out.exists()


def test_edits_are_debounced(client, config, monkeypatch):
    calls = []
    monkeypatch.setattr(publish, "publish_now", lambda: calls.append(1))
    for _ in range(5):
        publish.schedule(delay=0.2)
    publish._timer.join(2)
    assert calls == [1]
