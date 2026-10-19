"""Read-only Excel status report for execs, written to a shared (e.g. Box Drive) folder.

The tracker stays the source of truth: every publish builds a brand-new workbook from the case files
and swaps it onto `publish_path` (data/config.json). Nothing here may ever break a request.
"""
from __future__ import annotations

import logging
import os
import threading
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import engine, storage

log = logging.getLogger("contracts_tracker.publish")

DEBOUNCE_SECONDS = 10
SOON_DAYS = 30
HEADER_ROW = 3  # on the Active sheet: row 1 title, row 2 "Last updated"

ACTIVE_HEADERS = ["Case ID", "Type", "Store #", "Store Name", "Current Step", "Responsible Team",
                  "Days in Step", "Status", "Target Signing Date", "Progress", "Steps Done"]
ATTENTION_HEADERS = ["Case ID", "Type", "Store #", "Store Name", "Current Step", "Responsible Team", "Days in Step", "Status"]
SOON_HEADERS = ["Case ID", "Type", "Store #", "Store Name", "Target Signing Date", "Current Step",
                "Days to/Past Due", "Progress"]
COMPLETED_HEADERS = ["Case ID", "Type", "Store #", "Store Name", "Started", "Signed", "Completed"]
CLOSED_HEADERS = ["Case ID", "Type", "Store #", "Store Name", "Outcome", "Reason", "Started", "Closed"]

TYPE_ORDER = ["renewal", "transfer", "expansion", "lease"]
CLOSED_STATUSES = ("not_approved", "withdrawn")
OUTCOME = {"not_approved": "Not approved", "withdrawn": "Withdrawn"}

NAVY, GREY, STRIPE = "1F3864", "6B7785", "F2F5F9"
# Status key for execs: (fill, text colour, what it means), in order of precedence from the bottom up.
STATUS_STYLE = {
    "On Track": ("C6EFCE", "006100", "Progressing within the expected time frame for the current step."),
    "Delayed": ("FFEB9C", "9C5700", "The current step is taking longer than its expected time frame."),
    "On Hold": ("FCD5B4", "974706", "Paused pending actions or a decision outside the Contracts Team."),
    "Past Due": ("FFC7CE", "9C0006", "The target signing date has passed and the agreement is not yet signed."),
}
DATE_FMT = "mmm d, yyyy"
# Headings centered, all data left-justified, on every sheet.
HEAD_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_ALIGN = Alignment(horizontal="left", vertical="top")


# ---- data -----------------------------------------------------------------------------------

def _views() -> list[tuple[dict, dict]]:
    """(view, summary) per case, computed exactly the way GET /api/cases does."""
    cfg, out = storage.load_config(), []
    for c in storage.list_cases():
        tmpl = engine.TEMPLATES.get(c["type"])
        if not tmpl:
            continue
        engine.sync(c, tmpl)
        view = engine.build_view(c, tmpl, config=cfg)
        out.append((view, engine.summarize(view)))
    return out


def _date(iso: Optional[str]) -> Optional[date]:
    return engine.parse(iso[:10]) if iso else None


def _past_target(s: dict, t: date) -> bool:
    return bool(s["target_signing_date"]) and not s["signing_date"] and _date(s["target_signing_date"]) < t


def _status(s: dict, t: date) -> str:
    """One status per case, the most serious that applies: Past Due, then On Hold, then Delayed (the current
    step is past its time limit, the same flag the tracker shows), else On Track."""
    if _past_target(s, t):
        return "Past Due"
    if s.get("hold"):
        return "On Hold"
    if any(cur and cur["is_stale"] for cur in (s["current"], _alongside(s))):
        return "Delayed"
    return "On Track"


def _type_rank(s: dict) -> int:
    return TYPE_ORDER.index(s["type"]) if s["type"] in TYPE_ORDER else len(TYPE_ORDER)


def _alongside(s: dict) -> Optional[dict]:
    """The current step of a track that runs alongside the main one (the renewal inspection)."""
    side = s.get("side")
    return side["current"] if side and side["alongside"] else None


def _step(s: dict) -> str:
    side = s.get("side")
    main = s["current"]["label"] if s["current"] else (side["waiting_label"] if side and not side["done"] else "")
    other = _alongside(s)
    return f"{main} + {other['label']}" if main and other else main or (other or {}).get("label", "")


def _owner(s: dict) -> str:
    teams = [cur["owner"] for cur in (s["current"], _alongside(s)) if cur]
    return " + ".join(dict.fromkeys(teams))


# ---- formatting helpers ---------------------------------------------------------------------

def _sheet(ws, title: str, stamp: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    ws["A2"] = f"Last updated: {stamp}"
    ws["A2"].font = Font(italic=True, color=GREY)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text = f"Last updated: {stamp}"
    ws.oddFooter.right.text = "Page &P of &N"


def _heading(ws, row: int, text: str) -> int:
    ws.cell(row=row, column=1, value=text).font = Font(size=12, bold=True, color=NAVY)
    return row + 1


def _widen(ws, col: int, value) -> None:
    n = 12 if isinstance(value, date) else 6 if isinstance(value, float) else len(str(value or ""))
    dim = ws.column_dimensions[get_column_letter(col)]
    dim.width = min(max(dim.width or 0, n + 3), 50)


def _header(ws, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = HEAD_ALIGN
        _widen(ws, col, h)


def _table(ws, row: int, name: str, headers: list[str], rows: list[list], empty: str = "None right now.") -> int:
    """Write a banded Excel table with filter buttons at `row`; returns the next free row."""
    _header(ws, row, headers)
    if not rows:
        ws.cell(row=row + 1, column=1, value=empty).font = Font(italic=True, color=GREY)
        return row + 3
    for r, values in enumerate(rows, row + 1):
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.alignment = DATA_ALIGN
            h = headers[col - 1]
            if isinstance(v, date):
                cell.number_format = DATE_FMT
            if h == "Progress":
                cell.number_format = "0%"
            elif h == "Status" and v in STATUS_STYLE:
                fill, color, _ = STATUS_STYLE[v]
                cell.fill, cell.font = PatternFill("solid", fgColor=fill), Font(bold=True, color=color)
            elif h == "Reason":
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            _widen(ws, col, v)
    last = row + len(rows)
    table = Table(displayName=name, ref=f"A{row}:{get_column_letter(len(headers))}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    if "Progress" in headers:
        letter = get_column_letter(headers.index("Progress") + 1)
        ws.conditional_formatting.add(f"{letter}{row + 1}:{letter}{last}", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1, color="5B9BD5"))
    return last + 2


# ---- sheets ---------------------------------------------------------------------------------

def _store(s: dict) -> list:
    return [s["id"], s["type_label"], s["store_number"], s["store_name"]]


def _summary(ws, summaries: list[dict], views: list[dict], t: date, stamp: str) -> None:
    _sheet(ws, "Contracts Status", stamp)
    ws.sheet_properties.tabColor = NAVY
    active = [s for s in summaries if s["status"] == "active"]
    soon = sorted((s for s in active if not s["signing_date"] and s["target_signing_date"]
                   and _date(s["target_signing_date"]) <= t + timedelta(days=SOON_DAYS)),
                  key=lambda s: s["target_signing_date"])

    # counts by type
    row = _heading(ws, 4, "At a Glance")
    year = t.year
    headers = ["Case Type", "Active", *STATUS_STYLE, f"Due Within {SOON_DAYS} Days",
               f"Completed in {year}", "Closed Early"]
    _header(ws, row, headers)
    labels = {k: tm["label"] + "s" for k, tm in engine.TEMPLATES.items()}
    types = sorted(labels, key=lambda k: TYPE_ORDER.index(k) if k in TYPE_ORDER else len(TYPE_ORDER))
    totals = [0] * (len(headers) - 1)
    for i, k in enumerate(types):
        mine = [s for s in active if s["type"] == k]
        counts = [len(mine),  # one status per case, so the status columns add up to Active
                  *(sum(1 for s in mine if _status(s, t) == st) for st in STATUS_STYLE),
                  sum(1 for s in soon if s["type"] == k and not _past_target(s, t)),
                  sum(1 for v in views if v["type"] == k and v["status"] == "completed"
                      and (_date(v["completed_at"]) or t).year == year),
                  sum(1 for v in views if v["type"] == k and v["status"] in CLOSED_STATUSES
                      and _date(v["updated_at"]).year == year)]
        totals = [a + b for a, b in zip(totals, counts)]
        for col, v in enumerate([labels[k], *counts], 1):
            cell = ws.cell(row=row + 1 + i, column=col, value=v)
            cell.alignment = DATA_ALIGN
            if i % 2:
                cell.fill = PatternFill("solid", fgColor=STRIPE)
            _widen(ws, col, v)
    total_row = row + 1 + len(types)
    for col, v in enumerate(["Total", *totals], 1):
        cell = ws.cell(row=total_row, column=col, value=v)
        cell.font, cell.alignment = Font(bold=True), DATA_ALIGN
        cell.border = Border(top=Side(style="thin", color=NAVY))
    for r in range(row + 1, total_row + 1):  # make Delayed / On Hold / Past Due counts pop when they're not zero
        for st in ("Delayed", "On Hold", "Past Due"):
            col = headers.index(st) + 1
            if ws.cell(row=r, column=col).value:
                ws.cell(row=r, column=col).font = Font(bold=True, color=STATUS_STYLE[st][1])

    # what needs attention
    row = _heading(ws, total_row + 3, "Needs Attention")
    ws.cell(row=row, column=1, value="Delayed, on hold, or past due; longest in current step first").font = Font(italic=True, color=GREY)
    late = sorted((s for s in active if _status(s, t) != "On Track"),
                  key=lambda s: -(s["current"]["days_in_step"] if s["current"] else 0))
    row = _table(ws, row + 1, "NeedsAttention", ATTENTION_HEADERS, [
        [*_store(s), _step(s), _owner(s),
         s["current"]["days_in_step"] if s["current"] else None, _status(s, t)] for s in late],
        empty="Nothing is delayed, on hold, or past due.")

    # what's closing
    row = _heading(ws, row + 1, "Upcoming and Past Due Signings")
    ws.cell(row=row, column=1, value=f"Target signing date within the next {SOON_DAYS} days or already passed").font = Font(italic=True, color=GREY)
    soon_rows = []
    for s in soon:
        left = (_date(s["target_signing_date"]) - t).days
        when = "Today" if left == 0 else f"In {left} day{'s' if left != 1 else ''}" if left > 0 else f"{-left} day{'s' if left != -1 else ''} past due"
        soon_rows.append([*_store(s), _date(s["target_signing_date"]), _step(s), when, s["progress"]["percent"] / 100])
    first = row + 2
    row = _table(ws, row + 1, "SigningSoon", SOON_HEADERS, soon_rows,
                 empty=f"No target signing dates in the next {SOON_DAYS} days or past due.")
    for i, r in enumerate(soon_rows):
        if "past" in r[6]:
            ws.cell(row=first + i, column=7).font = Font(bold=True, color=STATUS_STYLE["Past Due"][1])

    # status key
    row = _heading(ws, row + 1, "Status Key")
    for i, (name, (fill, color, meaning)) in enumerate(STATUS_STYLE.items()):
        cell = ws.cell(row=row + i, column=1, value=name)
        cell.fill, cell.font, cell.alignment = PatternFill("solid", fgColor=fill), Font(bold=True, color=color), DATA_ALIGN
        ws.cell(row=row + i, column=2, value=meaning).font = Font(color=GREY)


def _active(ws, summaries: list[dict], t: date, stamp: str) -> None:
    _sheet(ws, "Active Cases", stamp)
    active = sorted((s for s in summaries if s["status"] == "active"),
                    key=lambda s: (_type_rank(s), s["target_signing_date"] or "9999", s["start_date"]))
    rows = []
    for s in active:
        cur, p = s["current"], s["progress"]
        rows.append([*_store(s), _step(s), _owner(s), cur["days_in_step"] if cur else None,
                     _status(s, t), _date(s["target_signing_date"]), p["percent"] / 100, f"{p['done']} of {p['total']}"])
    _table(ws, HEADER_ROW, "ActiveCases", ACTIVE_HEADERS, rows, empty="No active cases.")
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"


def _done(ws, views: list[dict], t: date, stamp: str) -> None:
    _sheet(ws, f"Completed & Closed in {t.year}", stamp)
    completed = sorted((v for v in views if v["status"] == "completed" and (_date(v["completed_at"]) or t).year == t.year),
                       key=lambda v: v["completed_at"] or "", reverse=True)
    closed = sorted((v for v in views if v["status"] in CLOSED_STATUSES and _date(v["updated_at"]).year == t.year),
                    key=lambda v: v["updated_at"], reverse=True)
    row = _heading(ws, 4, "Completed")
    row = _table(ws, row, "Completed", COMPLETED_HEADERS, [
        [*_store(v), _date(v["start_date"]), _date(v["signing_date"]), _date(v["completed_at"])] for v in completed],
        empty=f"None completed yet in {t.year}.")
    row = _heading(ws, row + 1, "Closed Early")
    _table(ws, row, "ClosedEarly", CLOSED_HEADERS, [
        [*_store(v), OUTCOME[v["status"]], v.get("closed_reason") or "", _date(v["start_date"]),
         _date(v["updated_at"])] for v in closed],
        empty=f"None closed early in {t.year}.")


def build_workbook(now: Optional[datetime] = None) -> Workbook:
    """Always a fresh workbook; an existing report file is never opened."""
    now = now or datetime.now()
    stamp = now.strftime("%b %d, %Y %I:%M %p").replace(" 0", " ")
    pairs = _views()
    views, summaries = [v for v, _ in pairs], [s for _, s in pairs]

    wb = Workbook()
    _summary(wb.active, summaries, views, now.date(), stamp)
    wb.active.title = "Summary"
    _active(wb.create_sheet("Active"), summaries, now.date(), stamp)
    _done(wb.create_sheet("Completed & Closed"), views, now.date(), stamp)
    wb.active = 0  # open on the Summary
    return wb



def _save(wb: Workbook, path: Path) -> None:
    """Write next to the target, then swap it in, so execs never see a half-written file."""
    tmp = path.with_name(f"~{path.stem}.{os.getpid()}.tmp.xlsx")
    wb.save(tmp)
    try:
        for attempt in range(5):  # Excel or Box Drive may hold the file for a moment
            try:
                os.replace(tmp, path)
                return
            except PermissionError:
                time.sleep(0.5 * (attempt + 1))
        os.replace(tmp, path)
    finally:
        if tmp.exists():
            tmp.unlink(missing_ok=True)


def publish_now() -> dict:
    """Build and save the report. Never raises; returns {"ok": bool, "detail": str}."""
    try:
        target = storage.load_config().get("publish_path")
        if not target:
            found = "has no publish_path" if storage.CONFIG_FILE.exists() else "doesn't exist"
            return {"ok": False, "detail": f"Publishing is off: {storage.CONFIG_FILE} {found}."}
        path = Path(target)
        if not path.parent.is_dir():
            raise FileNotFoundError(f"Folder not found: {path.parent}")
        with _PUBLISH_LOCK:
            _save(build_workbook(), path)
        log.info("Published report to %s", path)
        return {"ok": True, "detail": f"Published {path.name}."}
    except Exception as e:  # noqa: BLE001 - publishing must never take the app down
        log.exception("Publishing the report failed")
        return {"ok": False, "detail": f"Publish failed: {e}"}


# ---- debounced publishing after edits -------------------------------------------------------

_PUBLISH_LOCK = threading.Lock()
_TIMER_LOCK = threading.Lock()
_timer: Optional[threading.Timer] = None


def schedule(delay: float = DEBOUNCE_SECONDS) -> None:
    """Publish `delay` seconds after the last call, so a burst of edits makes one file."""
    global _timer
    try:
        with _TIMER_LOCK:
            if _timer:
                _timer.cancel()
            _timer = threading.Timer(delay, publish_now)
            _timer.daemon = True
            _timer.start()
    except Exception:  # noqa: BLE001
        log.exception("Could not schedule a publish")


def publish_in_background() -> None:
    threading.Thread(target=publish_now, daemon=True).start()
